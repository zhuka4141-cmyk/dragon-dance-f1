
# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image as _I
import os

WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1B,0x1B,0x1F); SUB=RGBColor(0x6B,0x6B,0x72)
RED=RGBColor(0xD9,0x1F,0x2A); GOLD=RGBColor(0xF2,0xA9,0x00); SOFT=RGBColor(0xF4,0xF4,0xF6)
HAIR=RGBColor(0xE3,0xE3,0xE8); LIGHT=RGBColor(0xEA,0xEA,0xEF); LINEDK=RGBColor(0x3A,0x3A,0x42)
CARD=RGBColor(0xFB,0xFB,0xFD); DARK=RGBColor(0x0D,0x0D,0x10); DARKCARD=RGBColor(0x20,0x20,0x26)
LATIN="Segoe UI"; EA="Microsoft JhengHei"
SRC="C:/Users/Administrator/dragon-dance-f1/assets/"
NORM="C:/tmp/dd/pptximg"; os.makedirs(NORM, exist_ok=True)
def norm(name):
    out=os.path.join(NORM,name); im=_I.open(SRC+name).convert("RGB"); im.save(out,"JPEG",quality=92); return out
for _n in ["race-car.jpg","manufacturing.jpg","team-moment-02.jpg","team-moment-01.jpg","sponsorship-marketing.jpg","dragon-mark.png"]:
    norm(_n)
A=NORM+"/"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]; M=Inches(0.7)

def _font(run,size,bold,color):
    f=run.font; f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name=LATIN
    rPr=run._r.get_or_add_rPr()
    for tag in ('a:latin','a:ea','a:cs'):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set('typeface', LATIN if tag=='a:latin' else EA)

def text(slide,x,y,w,h,lines,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.12):
    tb=slide.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p=tf.paragraphs[0]; p.alignment=align; p.line_spacing=ls
    for line in lines:
        if p is not tf.paragraphs[0]: p=tf.add_paragraph(); p.alignment=align; p.line_spacing=ls
        for t in line:
            r=p.add_run(); r.text=t[0]; _font(r,t[1],t[2] if len(t)>2 else False,t[3] if len(t)>3 else INK)
    return tb

def rect(slide,x,y,w,h,fill=None,shape=MSO_SHAPE.RECTANGLE,radius=None,alpha=None):
    sp=slide.shapes.add_shape(shape,x,y,w,h)
    if fill is None: sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb=fill
        if alpha is not None:
            srgb=sp.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            al=srgb.makeelement(qn('a:alpha'),{'val':str(int(alpha*1000))}); srgb.append(al)
    sp.line.fill.background(); sp.shadow.inherit=False
    if radius is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=radius
        except Exception: pass
    return sp

def hline(slide,x,y,w,color=HAIR,weight=1.0):
    ln=slide.shapes.add_connector(1,x,y,x+w,y); ln.line.color.rgb=color; ln.line.width=Pt(weight); ln.shadow.inherit=False

def pic(slide,path,x,y,w,h):
    """Add image cropped to fill the box (no distortion)."""
    im=slide.shapes.add_picture(path,x,y)
    im.width=w; im.height=h
    iw,ih=_I.open(path).size; src=iw/ih; box=(w/914400.0)/(h/914400.0)
    if src>box: frac=(1-box/src)/2; im.crop_left=frac; im.crop_right=frac
    else: frac=(1-src/box)/2; im.crop_top=frac; im.crop_bottom=frac
    return im

def footer(slide,idx,on_dark=False):
    c=LIGHT if on_dark else SUB
    text(slide,M,Inches(7.06),Inches(6),Inches(0.28),[[("TEAM DRAGON DANCE  ·  STEM RACING",9,False,c)]])
    text(slide,Inches(12.0),Inches(7.06),Inches(0.6),Inches(0.28),[[(f"{idx:02d}",9,True,GOLD if on_dark else RED)]],align=PP_ALIGN.RIGHT)

def kicker(slide,txt,line_color=HAIR,kc=RED):
    rect(slide,M,Inches(0.62),Inches(0.46),Inches(0.07),fill=kc)
    text(slide,Inches(1.32),Inches(0.45),Inches(9),Inches(0.4),[[(txt,12,True,kc)]])
    hline(slide,M,Inches(0.96),Inches(11.93),line_color)

# ===== SLIDE 1 COVER =====
s=prs.slides.add_slide(BLANK)
pic(s,A+"race-car.jpg",0,0,Inches(13.333),Inches(7.5))
rect(s,0,0,Inches(13.333),Inches(7.5),fill=DARK,alpha=74)
rect(s,0,Inches(7.28),Inches(13.333),Inches(0.22),fill=RED)
pic(s,A+"dragon-mark.png",M,Inches(0.5),Inches(0.9),Inches(0.9))
text(s,Inches(1.85),Inches(0.56),Inches(10),Inches(0.5),[[("STEM RACING  ·  HONG KONG & MACAU",14,True,GOLD)]])
text(s,M,Inches(2.1),Inches(11.8),Inches(1.9),[[("贊助計劃書",72,True,WHITE)]])
rect(s,M,Inches(3.98),Inches(1.35),Inches(0.1),fill=RED)
text(s,M,Inches(4.36),Inches(11.5),Inches(1.2),[[("您的支持，點燃一群年輕工程師的夢想，",19,False,LIGHT)],[("更將貴品牌與香港文化、教育、創新直接掛鉤。",19,False,LIGHT)]],ls=1.3)
text(s,M,Inches(6.0),Inches(11.5),Inches(0.6),[[("TEAM DRAGON DANCE",26,True,WHITE)]])
text(s,M,Inches(6.66),Inches(11.5),Inches(0.4),[[("AD&FD POHL Leung Sing Tak College",12,False,LIGHT)]])

# ===== SLIDE 2 CONTENTS =====
s=prs.slides.add_slide(BLANK)
rect(s,0,0,Inches(13.333),Inches(7.5),fill=WHITE)
kicker(s,"CONTENTS")
text(s,M,Inches(1.32),Inches(6),Inches(0.8),[[("目錄",40,True,INK)]])
items=[("01","我們與比賽","About the Competition"),("02","共同的願景","Our Vision"),("03","資金與宣傳","Funding & Promotion"),("04","贊助方案","Sponsorship Tiers"),("05","聯絡方法","Contact Us")]
y=Inches(2.5)
for num,zh,en in items:
    text(s,M,y,Inches(1.1),Inches(0.55),[[(num,24,True,RED)]])
    text(s,Inches(1.95),y,Inches(6),Inches(0.5),[[(zh,20,True,INK)]])
    text(s,Inches(1.95),y+Inches(0.4),Inches(6),Inches(0.35),[[(en,11,False,SUB)]])
    hline(s,M,y+Inches(0.8),Inches(6.6))
    y+=Inches(0.98)
px=Inches(8.0); py=Inches(1.45); pw=Inches(4.5)
rect(s,px,py,pw,Inches(5.5),fill=SOFT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
pic(s,A+"team-moment-02.jpg",px+Inches(0.25),py+Inches(0.25),pw-Inches(0.5),Inches(4.3))
text(s,px+Inches(0.3),py+Inches(4.75),pw-Inches(0.6),Inches(0.6),[[("我們在賽道上，也在文化裡。",13,True,INK)]])
footer(s,2)

# ===== SLIDE 3 SECTION 1 =====
s=prs.slides.add_slide(BLANK)
rect(s,0,0,Inches(13.333),Inches(7.5),fill=WHITE)
kicker(s,"01 · ABOUT")
text(s,M,Inches(1.28),Inches(6),Inches(0.8),[[("我們與比賽",34,True,INK)]])
text(s,M,Inches(2.02),Inches(6.4),Inches(1.1),[[("STEM RACING — F1 官方支持、全球最大型青少年工程競賽，讓 9 至 19 歲學生像真正的 F1 車隊一樣：設計、製造、競速、宣傳。",15,False,SUB)]],ls=1.35)
stats=[("180萬","青少年參賽"),("65","國家參與"),("20米","標準賽道"),("1秒","決勝時刻")]
sw=Inches(2.45); shh=Inches(1.12); sg=Inches(0.25); sy=Inches(3.5)
for i,(num,lab) in enumerate(stats):
    cx=M+(i%2)*(sw+sg); cy=sy+(i//2)*(shh+sg)
    rect(s,cx,cy,sw,shh,fill=SOFT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.16)
    text(s,cx+Inches(0.2),cy+Inches(0.14),sw-Inches(0.4),Inches(0.5),[[(num,24,True,RED)]])
    text(s,cx+Inches(0.2),cy+Inches(0.68),sw-Inches(0.4),Inches(0.35),[[(lab,11,False,SUB)]])
px=Inches(7.7); py=Inches(1.6); pw=Inches(5.0)
rect(s,px,py,pw,Inches(5.2),fill=SOFT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
pic(s,A+"manufacturing.jpg",px+Inches(0.25),py+Inches(0.25),pw-Inches(0.5),Inches(4.1))
text(s,px+Inches(0.3),py+Inches(4.55),pw-Inches(0.6),Inches(0.5),[[("一支紮根香港的隊伍，將舞火龍文化融入工程競賽。",12,True,INK)]])
footer(s,3)

# ===== SLIDE 4 SECTION 2 =====
s=prs.slides.add_slide(BLANK)
rect(s,0,0,Inches(13.333),Inches(7.5),fill=WHITE)
kicker(s,"02 · VISION")
text(s,M,Inches(1.28),Inches(6),Inches(0.8),[[("共同的願景",34,True,INK)]])
text(s,M,Inches(2.02),Inches(11),Inches(0.5),[[("傳承文化，驅動未來",16,True,RED)]])
cards=[("國際舞台","您的品牌，隨我們征戰國際，與全球頂尖隊伍同框，觸及 65 個國家、數百萬觀眾。"),
       ("文化價值","舞火龍是香港非物質文化遺產。您的支持，讓這份文化技藝在國際年輕一代眼前重生。"),
       ("社會價值","支持香港年輕一代以實力傳遞香港文化故事，讓品牌與教育、創新、文化保育緊緊相連。")]
cw=Inches(3.83); cg=Inches(0.22); cy=Inches(2.6); chh=Inches(2.9)
for i,(t,d) in enumerate(cards):
    cx=M+i*(cw+cg)
    rect(s,cx,cy,cw,chh,fill=CARD,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
    rect(s,cx,cy,Inches(0.13),chh,fill=RED)
    text(s,cx+Inches(0.35),cy+Inches(0.32),cw-Inches(0.7),Inches(0.55),[[(t,21,True,INK)]])
    text(s,cx+Inches(0.35),cy+Inches(1.0),cw-Inches(0.7),Inches(1.8),[[(d,14,False,SUB)]],ls=1.4)
pic(s,A+"sponsorship-marketing.jpg",M,Inches(5.75),Inches(11.93),Inches(1.55))
footer(s,4)

# ===== SLIDE 5 SECTION 3 =====
s=prs.slides.add_slide(BLANK)
rect(s,0,0,Inches(13.333),Inches(7.5),fill=WHITE)
kicker(s,"03 · FUNDING & PROMOTION")
text(s,M,Inches(1.28),Inches(6),Inches(0.8),[[("資金與宣傳",34,True,INK)]])
px=M; py=Inches(2.2); pw=Inches(4.2)
rect(s,px,py,pw,Inches(4.6),fill=SOFT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
pic(s,A+"race-car.jpg",px+Inches(0.22),py+Inches(0.22),pw-Inches(0.44),Inches(3.6))
text(s,px+Inches(0.26),py+Inches(4.0),pw-Inches(0.5),Inches(0.5),[[("線上線下同步曝光，最大化品牌推廣效益。",12,True,INK)]])
rx=Inches(5.35); rw=Inches(7.28)
rect(s,rx,Inches(2.2),rw,Inches(2.2),fill=SOFT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.06)
text(s,rx+Inches(0.3),Inches(2.4),rw-Inches(0.6),Inches(0.5),[[("資金",18,True,RED)]])
text(s,rx+Inches(0.3),Inches(2.98),rw-Inches(0.6),Inches(1.3),[[("車隊需自主籌措營運經費，涵蓋設計、製造、報名與行銷。您的支持將投入 3D 打印耗材、制服、宣傳物料與後勤。",14,False,SUB)]],ls=1.35)
rect(s,rx,Inches(4.6),rw,Inches(2.2),fill=SOFT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.06)
text(s,rx+Inches(0.3),Inches(4.8),rw-Inches(0.6),Inches(0.5),[[("宣傳",18,True,RED)]])
text(s,rx+Inches(0.3),Inches(5.38),rw-Inches(0.6),Inches(1.3),[[("透過社交媒體平台宣傳贊助商；在制服及賽車上展示貴公司標誌，於賽事現場與線上內容持續強化品牌識別。",14,False,SUB)]],ls=1.35)
footer(s,5)

# ===== SLIDE 6 SECTION 4 =====
s=prs.slides.add_slide(BLANK)
rect(s,0,0,Inches(13.333),Inches(7.5),fill=WHITE)
kicker(s,"04 · SPONSORSHIP TIERS")
text(s,M,Inches(1.28),Inches(6),Inches(0.8),[[("贊助方案",34,True,INK)]])
tiers=[("龍","2500","港元",["社交媒體感謝","標誌放制服和賽車"]),("龖","3000","港元",["龍 方案全部回報","標誌放商品上"]),("龘","3500","港元",["龖 方案全部回報","一條宣傳短片"])]
cw=Inches(3.83); cg=Inches(0.22); cy=Inches(2.45); chh=Inches(4.0)
for i,(nm,pr,unit,feats) in enumerate(tiers):
    cx=M+i*(cw+cg)
    if i==1:
        rect(s,cx,cy,cw,chh,fill=INK,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05); rect(s,cx,Inches(2.32),cw,Inches(0.36),fill=RED)
    else:
        rect(s,cx,cy,cw,chh,fill=CARD,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05); hline(s,cx+Inches(0.3),Inches(2.32),cw-Inches(0.6))
    nc=GOLD if i==1 else RED
    text(s,cx+Inches(0.32),cy+Inches(0.38),cw-Inches(0.64),Inches(0.7),[[(nm,38,True,nc)]])
    pc=WHITE if i==1 else INK; uc=LIGHT if i==1 else SUB
    text(s,cx+Inches(0.32),cy+Inches(1.28),cw-Inches(0.64),Inches(0.8),[[(pr,32,True,pc),("  "+unit,13,False,uc)]])
    yy=cy+Inches(2.45)
    for ft in feats:
        text(s,cx+Inches(0.32),yy,cw-Inches(0.64),Inches(0.6),[[("•  ",14,True,nc),(ft,14,False,INK if i!=1 else LIGHT)]],ls=1.1)
        yy+=Inches(0.62)
text(s,M,Inches(6.72),Inches(11.5),Inches(0.4),[[("條款細則歡迎隨時聯絡作查詢，表格供方案參考。",12,False,SUB)]])
footer(s,6)

# ===== SLIDE 7 CONTACT =====
s=prs.slides.add_slide(BLANK)
rect(s,0,0,Inches(13.333),Inches(7.5),fill=INK)
rect(s,0,Inches(7.28),Inches(13.333),Inches(0.22),fill=RED)
rect(s,M,Inches(0.62),Inches(0.46),Inches(0.07),fill=RED)
text(s,Inches(1.32),Inches(0.45),Inches(9),Inches(0.4),[[("05 · CONTACT US",12,True,GOLD)]])
hline(s,M,Inches(0.96),Inches(11.93),LINEDK)
px=M; py=Inches(1.5); pw=Inches(3.7)
rect(s,px,py,pw,Inches(5.0),fill=DARKCARD,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
pic(s,A+"dragon-mark.png",px+Inches(0.8),py+Inches(0.8),Inches(2.1),Inches(2.1))
text(s,px+Inches(0.3),py+Inches(3.4),pw-Inches(0.6),Inches(1.2),[[("TEAM",28,True,GOLD)],[("DRAGON DANCE",20,True,WHITE)]],ls=1.05)
text(s,Inches(4.8),Inches(1.32),Inches(8),Inches(0.8),[[("聯絡方法",38,True,WHITE)]])
text(s,Inches(4.8),Inches(2.2),Inches(8),Inches(0.5),[[("與我們展開對話，讓品牌與香港文化一起走向世界。",15,False,LIGHT)]])
contacts=[("EMAIL","dragondance@lstc.edu.hk"),("PHONE","+852 6011 6515"),("ADDRESS","Estate Secondary School No. 2, Shek Wai Kok Estate, Tsuen Wan, N.T., Hong Kong"),("SOCIAL MEDIA","@team_dragon_dance")]
y=Inches(3.05)
for lab,val in contacts:
    text(s,Inches(4.8),y,Inches(2.0),Inches(0.4),[[(lab,12,True,GOLD)]])
    text(s,Inches(6.8),y,Inches(6.0),Inches(0.5),[[(val,15,True,WHITE)]])
    hline(s,Inches(4.8),y+Inches(0.62),Inches(8.0),LINEDK)
    y+=Inches(0.95)
text(s,Inches(12.0),Inches(7.06),Inches(0.6),Inches(0.28),[[("07",9,True,GOLD)]],align=PP_ALIGN.RIGHT)

prs.save("C:/Users/Administrator/dragon-dance-f1/assets/sponsorship-proposal.pptx")
print("saved, slides:",len(prs.slides._sldIdLst))
